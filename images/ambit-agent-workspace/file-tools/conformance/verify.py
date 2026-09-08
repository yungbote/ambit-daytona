#!/usr/bin/env python3
"""Exercise installed tools as a non-root user with networking disabled."""
import importlib.metadata
import hashlib
import json
import os
from pathlib import Path
import re
import shutil
import signal
import subprocess
import sys
import tempfile
import time

from docx import Document
import pymupdf
from PIL import Image, ImageDraw


def run(argv, *, cwd, env=None, timeout=60):
    started = time.monotonic()
    result = subprocess.run(argv, cwd=cwd, env=env, capture_output=True,
                            text=True, timeout=timeout, check=False)
    if result.returncode:
        raise AssertionError({"argv": argv, "exit": result.returncode,
                              "stderr": result.stderr[-4000:]})
    return result.stdout, time.monotonic() - started


def measured_run(argv, *, cwd, timeout=60, cancel_after_output=None):
    """Linux wait4 gives this command's actual peak RSS, not a sampled estimate.

    Capture into files to avoid pipe backpressure. A deadline kills the command's
    process group, including any ordinary child tools, before reaping the owner.
    These are conformance measurements, not a new workspace execution policy.
    """
    started = time.monotonic()
    with tempfile.TemporaryFile() as stdout, tempfile.TemporaryFile() as stderr:
        process = subprocess.Popen(argv, cwd=cwd, stdout=stdout, stderr=stderr,
                                   start_new_session=True)
        cancelled = False
        cancellation = {}
        while True:
            waited, status, usage = os.wait4(process.pid, os.WNOHANG)
            if waited:
                break
            output_bytes = (cancel_after_output.stat().st_size
                            if cancel_after_output and cancel_after_output.exists() else 0)
            if output_bytes or time.monotonic() - started >= timeout:
                cancelled = True
                cancel_started = time.monotonic()
                try:
                    os.killpg(process.pid, signal.SIGKILL)
                except ProcessLookupError:
                    pass  # The command can finish between wait4 and the signal.
                _, status, usage = os.wait4(process.pid, 0)
                cancellation = {
                    "cancellation_trigger": "output_progress" if output_bytes else "deadline",
                    "observed_output_bytes": output_bytes,
                    "cancellation_seconds": time.monotonic() - cancel_started,
                }
                break
            time.sleep(0.01)
        process.returncode = os.waitstatus_to_exitcode(status)
        elapsed = time.monotonic() - started
        stdout.seek(0)
        stderr.seek(0)
        return {"argv": argv, "exit": process.returncode, "seconds": elapsed,
                "peak_rss_kib": usage.ru_maxrss, "cancelled": cancelled,
                **cancellation,
                "process_group": process.pid,
                "stdout": stdout.read().decode(errors="replace"),
                "stderr": stderr.read().decode(errors="replace")}


def command_evidence(result):
    return {key: value for key, value in result.items()
            if key not in {"stdout", "stderr", "process_group"}}


results = []
with tempfile.TemporaryDirectory(prefix="ambit-file-tools-") as name:
    directory = Path(name)
    originals = {}
    assert os.getuid() != 0, "Exercise the actual non-root workspace user"
    assert importlib.metadata.version("mammoth") == "1.12.1"
    private_python = "/opt/ambit/file-tools/python/bin/python"
    version, _ = run([private_python, "-c", "import importlib.metadata as m; print(m.version('markitdown')); print(m.version('mammoth'))"], cwd=directory)
    assert version.splitlines() == ["0.1.7", "1.11.0"]
    results.append({"check": "isolated Python dependencies", "passed": True})

    doc = Document()
    doc.add_heading("Workspace quantity check", 1)
    doc.add_paragraph("The source material remains unchanged.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Material"
    table.rows[0].cells[1].text = "Quantity"
    row = table.add_row()
    row.cells[0].text = "Copper"
    row.cells[1].text = "18.75"
    source = directory / "source document.docx"
    doc.save(source)
    originals[source] = hashlib.sha256(source.read_bytes()).hexdigest()
    markdown, elapsed = run(["markitdown", source.name], cwd=directory)
    assert re.search(r"(?m)^#\s+Workspace quantity check\s*$", markdown), markdown
    assert re.search(r"(?m)^\|\s*Material\s*\|\s*Quantity\s*\|\s*$", markdown), markdown
    assert re.search(r"(?m)^\|\s*Copper\s*\|\s*18\.75\s*\|\s*$", markdown), markdown
    results.append({"check": "MarkItDown structure", "seconds": elapsed, "passed": True})

    # Detection comes from content; this is not an application extension router.
    opaque = directory / "opaque.upload"
    opaque.write_bytes(source.read_bytes())
    originals[opaque] = hashlib.sha256(opaque.read_bytes()).hexdigest()
    text, elapsed = run(["tika", "--text", opaque.name], cwd=directory)
    assert all(value in text for value in ["Workspace quantity check", "Copper", "18.75"])
    results.append({"check": "Tika content detection and relative path", "seconds": elapsed, "passed": True})

    image = Image.new("RGB", (1800, 1200), "white")
    ImageDraw.Draw(image).text((100, 100), "RASTER EVIDENCE 427", fill="black", font_size=60)
    image_path = directory / "image.png"
    image.save(image_path)
    originals[image_path] = hashlib.sha256(image_path.read_bytes()).hexdigest()
    pdf = pymupdf.open()
    page = pdf.new_page(width=900, height=600)
    page.insert_image(page.rect, filename=str(image_path))
    pdf_path = directory / "image-only.pdf"
    pdf.save(pdf_path)
    pdf.close()
    originals[pdf_path] = hashlib.sha256(pdf_path.read_bytes()).hexdigest()
    traps = directory / "traps"
    traps.mkdir()
    trap = traps / "tesseract"
    marker = directory / "tesseract-invocations.jsonl"
    real_tesseract = shutil.which("tesseract")
    assert real_tesseract
    # Tika 4.0.0 initializes TesseractOCRParser by invoking tesseract with no
    # arguments (hasTesseract), even when skipOcr disables extraction. Record
    # exact argv and delegate unchanged: an unavailable fake OCR binary would
    # make this negative test pass for the wrong reason.
    # https://github.com/apache/tika/blob/4.0.0/tika-parsers/tika-parsers-standard/tika-parsers-standard-modules/tika-parser-ocr-module/src/main/java/org/apache/tika/parser/ocr/TesseractOCRParser.java
    trap.write_text(f"#!{sys.executable}\nimport json, os, sys\n"
                    "with open(os.environ['OCR_PROBE_MARKER'], 'a') as f:\n"
                    "    f.write(json.dumps(sys.argv[1:]) + '\\n')\n"
                    "os.execv(os.environ['OCR_PROBE_EXECUTABLE'], "
                    "[os.environ['OCR_PROBE_EXECUTABLE'], *sys.argv[1:]])\n")
    trap.chmod(0o755)
    environment = {**os.environ, "PATH": str(traps) + ":" + os.environ["PATH"],
                   "OCR_PROBE_MARKER": str(marker),
                   "OCR_PROBE_EXECUTABLE": real_tesseract}
    # Include an image itself: PDF NO_OCR alone does not disable the image parser.
    for raster in (pdf_path, image_path):
        marker.unlink(missing_ok=True)
        text, elapsed = run(["tika", "--text", raster.name], cwd=directory,
                            env=environment, timeout=30)
        calls = [json.loads(line) for line in marker.read_text().splitlines()]
        assert calls, "The availability probe must actually exercise the spy"
        assert all(args == [] for args in calls), {"unexpected OCR argv": calls}
        assert "RASTER EVIDENCE" not in text
        results.append({"check": "default extraction performs no OCR",
                        "input": raster.name, "seconds": elapsed,
                        "tesseract_argv": calls, "passed": True})

    marker.unlink()
    text, elapsed = run(["tesseract", image_path.name, "stdout", "--psm", "6"],
                        cwd=directory, env=environment, timeout=30)
    calls = [json.loads(line) for line in marker.read_text().splitlines()]
    assert calls == [[image_path.name, "stdout", "--psm", "6"]]
    assert "RASTER EVIDENCE 427" in text, text
    results.append({"check": "explicit OCR remains usable and is observed",
                    "seconds": elapsed, "tesseract_argv": calls,
                    "text": text.strip(), "passed": True})

    explicit_config = directory / "explicit-ocr.json"
    explicit_config.write_text(json.dumps({"parsers": [
        {"pdf-parser": {"ocr": {"strategy": "OCR_ONLY"}}},
        {"tesseract-ocr-parser": {"skipOcr": False}},
        {"default-parser": {}},
    ]}))
    marker.unlink()
    text, elapsed = run(["tika", "--config=" + explicit_config.name, "--text", pdf_path.name],
                        cwd=directory, env=environment, timeout=30)
    calls = [json.loads(line) for line in marker.read_text().splitlines()]
    assert any(args and args[0] not in {"--version", "--list-langs"}
               for args in calls), calls
    assert "RASTER EVIDENCE 427" in text, text
    results.append({"check": "explicit Tika config can enable OCR",
                    "seconds": elapsed, "tesseract_argv": calls,
                    "text": text.strip(), "passed": True})

    version, _ = run(["vips", "--version"], cwd=directory)
    assert "8.16.1" in version
    run(["vips", "crop", str(image_path), "region.png", "80", "80", "1000", "200"], cwd=directory)
    with Image.open(directory / "region.png") as region:
        assert region.size == (1000, 200)
        assert region.tobytes() == image.crop((80, 80, 1080, 280)).tobytes()
    results.append({"check": "exact image region", "passed": True})

    # 160 MP tiled input exceeds the host vision limit. General file tools must
    # still be able to inspect a small chosen region without a full-image copy.
    run(["vips", "black", "large.v", "16000", "10000", "--bands", "3"], cwd=directory)
    run(["vips", "draw_rect", "large.v", "24 96 168", "12340", "8765", "720", "360", "--fill"], cwd=directory)
    run(["vips", "tiffsave", "large.v", "large.tiff", "--tile", "--compression", "deflate"], cwd=directory)
    (directory / "large.v").unlink()
    large_hash = hashlib.sha256((directory / "large.tiff").read_bytes()).hexdigest()
    region_result = measured_run(["vips", "crop", "large.tiff", "large-region.png", "12280", "8725", "800", "450"], cwd=directory)
    assert region_result["exit"] == 0 and not region_result["cancelled"], region_result
    # A regression bound for this exact tiled fixture, not a universal RSS SLA.
    assert region_result["peak_rss_kib"] < 256 * 1024, region_result
    expected = Image.new("RGB", (800, 450), "black")
    ImageDraw.Draw(expected).rectangle((60, 40, 779, 399), fill=(24, 96, 168))
    with Image.open(directory / "large-region.png") as region:
        assert region.size == expected.size
        assert region.convert("RGB").tobytes() == expected.tobytes()
    assert hashlib.sha256((directory / "large.tiff").read_bytes()).hexdigest() == large_hash
    results.append({"check": "large tiled input exact region and peak RSS",
                    "source_pixels": 160_000_000,
                    "region_pixels": 800 * 450,
                    **command_evidence(region_result), "passed": True})

    malformed = directory / "broken.png"
    malformed.write_bytes(image_path.read_bytes()[:40])
    bad_image = measured_run(["vips", "copy", malformed.name, "broken-output.png"], cwd=directory, timeout=10)
    assert bad_image["exit"] != 0 and not bad_image["cancelled"], bad_image
    assert bad_image["stderr"].strip(), bad_image
    broken_pdf = directory / "broken.pdf"
    broken_pdf.write_bytes(b"%PDF-1.7\n1 0 obj\n<< /Type /Catalog >>\n")
    bad_pdf = measured_run(["qpdf", "--check", broken_pdf.name], cwd=directory, timeout=10)
    assert bad_pdf["exit"] != 0 and not bad_pdf["cancelled"], bad_pdf
    results.append({"check": "malformed image and PDF return bounded errors",
                    "commands": [command_evidence(bad_image), command_evidence(bad_pdf)],
                    "passed": True})

    partial = directory / "cancelled.partial.tiff"
    cancelled = measured_run(["vips", "gaussnoise", partial.name, "16000", "10000"],
                             cwd=directory, timeout=10, cancel_after_output=partial)
    assert cancelled["cancelled"] and cancelled["exit"] == -signal.SIGKILL, cancelled
    assert cancelled["cancellation_trigger"] == "output_progress", cancelled
    assert cancelled["observed_output_bytes"] > 0, cancelled
    assert cancelled["cancellation_seconds"] < 3, cancelled
    try:
        os.killpg(cancelled["process_group"], 0)
    except ProcessLookupError:
        pass
    else:
        raise AssertionError("Cancelled tool process group remains alive")
    # Partial output is explicitly uncommitted; killing a general CLI need not
    # remove its files. The ordinary Run/artifact owner decides publication.
    partial_bytes = partial.stat().st_size if partial.exists() else 0
    partial.unlink(missing_ok=True)
    run(["vips", "crop", str(image_path), "after-cancel.png", "80", "80", "1000", "200"], cwd=directory)
    with Image.open(directory / "after-cancel.png") as recovered:
        assert recovered.tobytes() == image.crop((80, 80, 1080, 280)).tobytes()
    results.append({"check": "active tool cancellation and subsequent work",
                    **command_evidence(cancelled), "partial_bytes": partial_bytes,
                    "passed": True})

    assert all(hashlib.sha256(path.read_bytes()).hexdigest() == digest
               for path, digest in originals.items())
    results.append({"check": "original files unchanged", "passed": True})

print(json.dumps({"checks": results, "passed": len(results)}, indent=2))
